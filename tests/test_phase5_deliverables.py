"""Fase 5 — Excel framework, paket Power BI, executive report, README.

Acceptance: semua angka di report dan Excel sama dengan data/output; data/output sendiri identik dengan
hasil pipeline yang dijalankan ulang (seed 42)."""
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pdfplumber
import pytest

from erm import datastore, export_powerbi, scoring
from erm.__main__ import run_engine
from erm.generator import generate_all

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "excel" / "ERM_Framework.xlsx"
PDF = ROOT / "report" / "ERM_Executive_Report.pdf"
FACTS = ROOT / "report" / "ERM_Executive_Report_facts.json"
MODEL = ROOT / "dashboard" / "model"
SHEETS = ["01_Risk_Taxonomy", "02_Risk_Appetite_Statement", "03_KRI_Definitions", "04_Thresholds",
          "05_Scoring_Methodology", "06_Risk_Assessment", "07_Stress_Scenarios", "08_Management_Actions",
          "09_Regulatory_Mapping", "10_Change_Log"]
REP = "2026-09"
TOL = 1e-9


@pytest.fixture(scope="module")
def out():
    return datastore.load_outputs()


@pytest.fixture(scope="module")
def ref():
    return datastore.load_reference()


# ---------- data/output adalah hasil pipeline terkini ----------
def test_data_output_matches_fresh_pipeline(tmp_path, out):
    raw, o = tmp_path / "raw", tmp_path / "out"
    generate_all(raw)
    run_engine(raw, o)
    for name, df in out.items():
        pd.testing.assert_frame_equal(pd.read_csv(o / f"{name}.csv"), df, check_exact=False, rtol=1e-12, obj=name)


# ---------- Excel ----------
@pytest.fixture(scope="module")
def wb_values():
    assert XLSX.exists(), "jalankan `make report`"
    return openpyxl.load_workbook(XLSX, data_only=True)


@pytest.fixture(scope="module")
def wb_formulas():
    return openpyxl.load_workbook(XLSX)


@pytest.fixture(scope="module")
def layout():
    return json.loads((ROOT / "excel" / "ERM_Framework_layout.json").read_text())


def test_excel_has_10_sheets_in_spec_order(wb_values):
    assert wb_values.sheetnames == SHEETS


def test_excel_recalculated_without_errors(wb_values, wb_formulas):
    n_formula, errors, missing = 0, [], []
    for ws in wb_formulas.worksheets:
        wv = wb_values[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_formula += 1
                    v = wv[c.coordinate].value
                    if v is None:
                        missing.append(f"{ws.title}!{c.coordinate}")
                    elif isinstance(v, str) and v.startswith("#"):
                        errors.append(f"{ws.title}!{c.coordinate}={v}")
    assert n_formula > 500 and not errors and not missing[:5]


def test_sheet05_credit_example_uses_live_formulas(wb_formulas, wb_values, layout, out):
    ws, wv = wb_formulas["05_Scoring_Methodology"], wb_values["05_Scoring_Methodology"]
    r = layout["credit_final_row"]
    for rr in range(r - 6, r + 3):
        assert str(ws.cell(rr, 2).value).startswith("="), rr
    score_cells = [ws.cell(rr, 9).value for rr in range(32, 37)]
    assert all(str(v).startswith("=IF(") for v in score_cells)
    engine = out["risk_scores_monthly"].query("period == @REP and risk_type == 'Credit'").final_score.iloc[0]
    assert abs(wv.cell(r, 2).value - engine) < TOL
    assert wv.cell(r + 2, 2).value == scoring.rating(engine)[0]


def test_sheet06_formulas_reconcile_to_engine(wb_values, layout, out):
    ws = wb_values["06_Risk_Assessment"]
    a0, a1 = layout["kri_rows"]
    km = out["kri_monthly"].query("period == @REP").set_index("kri_id")
    for r in range(a0, a1 + 1):
        kid = ws.cell(r, 1).value
        assert abs(ws.cell(r, 11).value - km.at[kid, "value"]) < TOL
        assert ws.cell(r, 13).value == km.at[kid, "score"] == ws.cell(r, 12).value, kid
        assert abs(ws.cell(r, 15).value - km.at[kid, "utilization"]) < 1e-9, kid
    rs = out["risk_scores_monthly"].query("period == @REP").set_index("risk_type")
    b0, b1 = layout["risk_rows"]
    for r in range(b0, b1 + 1):
        risk = ws.cell(r, 1).value
        assert abs(ws.cell(r, 3).value - rs.at[risk, "weighted_score"]) < TOL, risk
        assert abs(ws.cell(r, 6).value - rs.at[risk, "model_score"]) < TOL, risk
        assert abs(ws.cell(r, 8).value - rs.at[risk, "final_score"]) < TOL, risk
        assert ws.cell(r, 9).value == scoring.r1(rs.at[risk, "final_score"]) and ws.cell(r, 10).value == rs.at[risk, "rating"]
    ep = out["enterprise_profile_monthly"].query("period == @REP").iloc[0]
    c0 = layout["c0"]
    vals = [ws.cell(c0 + i, 2).value for i in range(10)]
    assert abs(vals[3] - ep.enterprise_score) < TOL and vals[4] == scoring.r1(ep.enterprise_score) and vals[5] == ep.rating
    assert abs(vals[6] - ep.limit_utilization) < 1e-9
    assert (vals[7], vals[8], vals[9]) == (ep.tolerance_breaches, ep.limit_breaches, ep.emerging_count)


def test_sheet04_thresholds_equal_reference(wb_values, ref):
    ws = wb_values["04_Thresholds"]
    ra = ref["risk_appetite"]
    for i, row in enumerate(ra.itertuples(index=False)):
        r = 5 + i
        assert [ws.cell(r, j + 1).value for j in range(len(ra.columns))] == [v.item() if hasattr(v, "item") else v for v in row]
        assert ws.cell(r, len(ra.columns) + 1).value is True
        if row.risk_type != "Capital Overlay":
            assert abs(ws.cell(r, len(ra.columns) + 2).value - 1.0) < 1e-9


def test_sheet07_stress_and_capital_equal_output(wb_values, layout, out):
    ws = wb_values["07_Stress_Scenarios"]
    cap = out["capital_overlay"].set_index("scenario")
    kr = layout["capital_rows"]
    for j, s in enumerate(datastore.SCENARIOS):
        assert abs(ws.cell(kr["kpmm"], 2 + j).value - cap.at[s, "kpmm"]) < 1e-9, s
        assert abs(ws.cell(kr["total"], 2 + j).value - cap.at[s, "total_loss"]) < 1e-9, s
        assert ws.cell(kr["score"], 2 + j).value == cap.at[s, "cap_score"]
    srs = out["stress_risk_scores"].pivot(index="risk_type", columns="scenario", values="final_score")
    se = out["stress_enterprise_summary"].set_index("scenario")
    er = layout["enterprise_stress_row"]
    for r in range(er - 10, er):
        risk = ws.cell(r, 1).value
        for j, s in enumerate(datastore.SCENARIOS):
            assert abs(ws.cell(r, 2 + j).value - srs.at[risk, s]) < TOL
    for j, s in enumerate(datastore.SCENARIOS):
        assert abs(ws.cell(er, 2 + j).value - se.at[s, "enterprise_score"]) < TOL
        assert ws.cell(er + 1, 2 + j).value == se.at[s, "rating"]
        assert (ws.cell(er + 3, 2 + j).value, ws.cell(er + 4, 2 + j).value) == (se.at[s, "tolerance_breaches"], se.at[s, "limit_breaches"])


def test_sheet08_actions_equal_output(wb_values, out):
    ws = wb_values["08_Management_Actions"]
    act = out["management_actions"]
    for i, row in enumerate(act.itertuples(index=False)):
        assert [ws.cell(12 + i, j + 1).value for j in range(len(act.columns))] == list(row)
    counts = act.status.value_counts()
    assert ws["B5"].value == counts.get("Open", 0) and ws["B6"].value == counts.get("Overdue", 0)
    assert ws["B9"].value == len(act)


# ---------- Power BI ----------
def test_powerbi_model_files_current_and_star_schema(out, ref):
    model = export_powerbi.build_model(out, ref)
    expected = {"dim_date", "dim_risk", "dim_kri", "dim_scenario", "dim_business_unit", "fact_kri_monthly",
                "fact_risk_score", "fact_enterprise", "fact_stress", "fact_capital", "fact_actions", "fact_emerging",
                "fact_interdependency"}
    assert set(model) == expected == {p.stem for p in MODEL.glob("*.csv")}
    for name, df in model.items():
        disk = pd.read_csv(MODEL / f"{name}.csv")
        assert list(disk.columns) == list(df.columns) and len(disk) == len(df), name
    fk = {"date_key": "dim_date", "risk_key": "dim_risk", "kri_key": "dim_kri", "scenario_key": "dim_scenario", "bu_key": "dim_business_unit"}
    for name, df in model.items():
        if not name.startswith("fact"):
            continue
        for col, dim in fk.items():
            if col in df.columns:
                assert set(df[col].dropna()) <= set(model[dim][col]), (name, col)
    fe = model["fact_enterprise"].query("date_key == 202609 and scenario_key == 0").iloc[0]
    ep = out["enterprise_profile_monthly"].query("period == @REP").iloc[0]
    assert abs(fe.enterprise_score - ep.enterprise_score) < TOL and fe.tolerance_breaches == ep.tolerance_breaches
    assert not list(ROOT.rglob("*.pbix"))


def test_dax_measures_cover_spec_and_reference_model_columns():
    text = (ROOT / "dashboard" / "DAX_measures.md").read_text(encoding="utf-8")
    for m in ["Enterprise Score", "Enterprise Rating", "Limit Utilization %", "Tolerance Breaches", "Limit Breaches",
              "KRI Score MoM Δ", "Deteriorating KRI Count", "Stress Δ vs Baseline", "KPMM Stressed", "Overdue Actions"]:
        assert re.search(rf"^### \d+\. {re.escape(m)}$", text, re.M), m
        assert f"{m} =" in text, m
    cols = {p.stem: set(pd.read_csv(p, nrows=0).columns) for p in MODEL.glob("*.csv")}
    refs = set(re.findall(r"\b(dim_\w+|fact_\w+)\[(\w+)\]", text))
    assert refs
    bad = [(t, c) for t, c in refs if c not in cols.get(t, set())]
    assert not bad, bad


def test_powerbi_spec_layout_pages():
    text = (ROOT / "dashboard" / "powerbi_spec.md").read_text(encoding="utf-8")
    for i, page in enumerate(["Enterprise Overview", "Risk Appetite", "Risk Trend & Early Warning", "Enterprise Stress",
                              "Interdependency", "Management Action"], start=1):
        assert f"## Halaman {i} — {page}" in text
    for hexcode in ("#2E7D32", "#F9A825", "#C62828"):
        assert hexcode in text
    cols = {p.stem: set(pd.read_csv(p, nrows=0).columns) for p in MODEL.glob("*.csv")}
    bad = [(t, c) for t, c in set(re.findall(r"`(dim_\w+|fact_\w+)\[(\w+)\]`", text)) if c not in cols.get(t, set())]
    assert not bad, bad


# ---------- Report ----------
@pytest.fixture(scope="module")
def pdf_text():
    assert PDF.exists(), "jalankan `make report`"
    with pdfplumber.open(PDF) as pdf:
        pages = [p.extract_text() or "" for p in pdf.pages]
    return pages


def test_report_page_count_and_sections(pdf_text):
    assert 8 <= len(pdf_text) <= 10
    full = "\n".join(pdf_text)
    for title in ["Executive Summary", "Framework", "Risk Profile", "Risk Appetite Monitoring", "Key Developments",
                  "Stress Testing", "Capital Overlay", "Risk Interdependency", "Emerging Risks", "Management Actions",
                  "Methodology, Limitations & Disclaimer"]:
        assert title in full, title
    assert "fully synthetic data" in re.sub(r"\s+", " ", full)


def _resolve(source: str, out: dict, ref: dict):
    """Nilai acuan untuk string sumber yang dicatat report (dihitung ulang dari data/output)."""
    def period(tok):
        return REP if tok == "rep" else tok
    m = re.fullmatch(r"kri_monthly\.(value|utilization)@([\w-]+):(\w+)", source)
    if m:
        k = out["kri_monthly"]
        return k[(k.period == period(m.group(2))) & (k.kri_id == m.group(3))][m.group(1)].iloc[0]
    m = re.fullmatch(r"kri_monthly\.trend_label=Deteriorating@rep", source)
    if m:
        k = out["kri_monthly"].query("period == @REP")
        return int(((k.trend_label == "Deteriorating") & (k.risk_type != "Capital Overlay")).sum())
    if source == "kri_monthly.early_warning_flag@rep":
        return int(out["kri_monthly"].query("period == @REP").early_warning_flag.astype(bool).sum())
    m = re.fullmatch(r"risk_scores_monthly\.(\w+)@([\w-]+):(.+)", source)
    if m:
        r = out["risk_scores_monthly"]
        return r[(r.period == period(m.group(2))) & (r.risk_type == m.group(3))][m.group(1)].iloc[0]
    m = re.fullmatch(r"enterprise_profile_monthly\.(\w+)(?:/\w+)?@([\w-]+)", source)
    if m:
        e = out["enterprise_profile_monthly"]
        return e[e.period == period(m.group(2))][m.group(1)].iloc[0]
    m = re.fullmatch(r"stress_risk_scores\.final_score:(\w+):(.+)", source)
    if m:
        r = out["stress_risk_scores"]
        return r[(r.scenario == m.group(1)) & (r.risk_type == m.group(2))].final_score.iloc[0]
    m = re.fullmatch(r"stress_enterprise_summary\.(\w+)(?:/\w+)?:(\w+)", source)
    if m:
        return out["stress_enterprise_summary"].set_index("scenario").at[m.group(2), m.group(1)]
    m = re.fullmatch(r"capital_overlay\.(\w+):(\w+)", source)
    if m:
        return out["capital_overlay"].set_index("scenario").at[m.group(2), m.group(1)]
    if source == "capital_overlay.max(loss)/total_loss:severe":
        c = out["capital_overlay"].set_index("scenario").loc["severe"]
        return max(c.credit_loss, c.investment_loss, c.op_loss, c.dcr_cost) / c.total_loss
    m = re.fullmatch(r"risk_weights\.materiality_weight:(.+)", source)
    if m:
        return ref["risk_weights"].set_index("risk_type").at[m.group(1), "materiality_weight"]
    if source == "risk_weights.materiality_weight.sum":
        return ref["risk_weights"].materiality_weight.sum()
    if source == "risk_appetite (tanpa Capital Overlay)":
        return int((ref["risk_appetite"].risk_type != "Capital Overlay").sum())
    m = re.fullmatch(r"emerging_risks\.priority_score:(\w+)", source)
    if m:
        return ref["emerging_risks"].set_index("er_id").at[m.group(1), "priority_score"]
    a = out["management_actions"]
    m = re.fullmatch(r"management_actions\.status=(\w+)", source)
    if m:
        return int((a.status == m.group(1)).sum())
    if source == "management_actions.status in Open/Overdue":
        return int(a.status.isin(["Open", "Overdue"]).sum())
    if source == "management_actions.count":
        return len(a)
    raise KeyError(source)


def test_report_every_number_traced_to_outputs(out, ref, pdf_text):
    facts = json.loads(FACTS.read_text(encoding="utf-8"))
    assert len(facts) > 100
    squashed = re.sub(r"\s+", "", "".join(pdf_text))
    unresolved, wrong, absent = [], [], []
    for f in facts:
        try:
            expected = _resolve(f["source"], out, ref)
        except KeyError:
            unresolved.append(f["source"])
            continue
        if isinstance(f["value"], str):
            if f["value"] != expected:
                wrong.append((f["key"], f["value"], expected))
        elif abs(float(f["value"]) - float(expected)) > 1e-9:
            wrong.append((f["key"], f["value"], expected))
        if re.sub(r"\s+", "", f["text"]) not in squashed:
            absent.append((f["key"], f["text"]))
    assert not unresolved, unresolved
    assert not wrong, wrong
    assert not absent, absent


def test_report_headline_numbers_match_outputs(out, pdf_text):
    page1 = re.sub(r"\s+", " ", pdf_text[0])
    ep = out["enterprise_profile_monthly"].query("period == @REP").iloc[0]
    cap = out["capital_overlay"].set_index("scenario")
    se = out["stress_enterprise_summary"].set_index("scenario")
    fmt = lambda x: f"{x:.1f}".replace(".", ",")
    for s in [fmt(scoring.r1(ep.enterprise_score)), ep.rating, f"{round(ep.limit_utilization)}%",
              f"{ep.tolerance_breaches} / {ep.limit_breaches}", fmt(cap.at["severe", "kpmm"]) + "%",
              fmt(scoring.r1(se.at["severe", "enterprise_score"]))]:
        assert s in page1, s


# ---------- README ----------
def test_readme_sections():
    text = (ROOT / "README.md").read_text(encoding="utf-8")
    for h in ["## Cara menjalankan", "## Arsitektur", "## Keputusan metodologi", "## Limitation", "fully synthetic data"]:
        assert h in text, h
    for target in ["make data", "make engine", "make report", "make test", "make all"]:
        assert target in text, target
