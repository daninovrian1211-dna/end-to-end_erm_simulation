"""Penutupan project — struktur final spec §13, notebook (hanya memanggil src/), regulatory mapping,
dan konsistensi tabel hasil README dengan data/output."""
import json
import os
import re
from pathlib import Path

import matplotlib
import openpyxl
import pytest

from erm import datastore, scoring

matplotlib.use("Agg")
ROOT = Path(__file__).resolve().parents[1]
NOTEBOOKS = ROOT / "notebooks"


def test_final_repository_structure_spec_13():
    required = [
        "README.md", "Makefile", "config/bank_profile.yaml", "config/stress_parameters.yaml",
        "data/reference/risk_appetite.csv", "data/reference/risk_weights.csv", "data/reference/macro_scenarios.csv",
        "data/reference/judgment_overrides.csv", "data/reference/interdependency_edges.csv",
        "data/reference/emerging_risks.csv", "data/reference/action_library.csv",
        "src/erm/generator/calibrate.py", "src/erm/generator/balance_sheet.py", "src/erm/generator/credit.py",
        "src/erm/generator/liquidity.py", "src/erm/validation.py", "src/erm/kri.py", "src/erm/scoring.py",
        "src/erm/stress.py", "src/erm/capital.py", "src/erm/actions.py", "src/erm/export_powerbi.py",
        "excel/ERM_Framework.xlsx", "dashboard/powerbi_spec.md", "dashboard/DAX_measures.md",
        "regulatory/Regulatory_Mapping.xlsx", "report/ERM_Executive_Report.pdf",
    ]
    missing = [p for p in required if not (ROOT / p).exists()]
    assert not missing, missing
    assert len(list((ROOT / "data" / "raw").glob("*.csv"))) == 15
    assert len(list((ROOT / "data" / "output").glob("*.csv"))) == len(datastore.OUTPUT_FILES)


def _code_cells(path: Path) -> list[str]:
    nb = json.loads(path.read_text(encoding="utf-8"))
    assert nb["nbformat"] == 4
    return ["".join(c["source"]) for c in nb["cells"] if c["cell_type"] == "code"]


def test_notebooks_01_to_05_only_call_src():
    books = sorted(NOTEBOOKS.glob("*.ipynb"))
    assert [b.name[:2] for b in books] == ["01", "02", "03", "04", "05"]
    for b in books:
        code = "\n".join(_code_cells(b))
        assert "from erm import" in code, b.name
        assert not re.search(r"^\s*(def|class)\s", code, re.M), f"{b.name}: logika harus di src/erm"


@pytest.mark.parametrize("name", ["01_data_generation.ipynb", "02_kri_risk_appetite.ipynb", "03_trend_early_warning.ipynb",
                                  "04_stress_capital.ipynb", "05_reporting.ipynb"])
def test_notebooks_execute(name, monkeypatch):
    monkeypatch.chdir(NOTEBOOKS)
    ns: dict = {}
    for cell in _code_cells(NOTEBOOKS / name):
        exec(compile(cell, name, "exec"), ns)
    matplotlib.pyplot.close("all")


def test_regulatory_mapping_verified_status():
    from erm.excel_framework import REG_MAPPING, REG_SOURCES, REG_VERIFIED_ON
    for path, sheet in ((ROOT / "regulatory" / "Regulatory_Mapping.xlsx", None), (ROOT / "excel" / "ERM_Framework.xlsx", "09_Regulatory_Mapping")):
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet] if sheet else wb.active
        assert ws.cell(4, 3).value == f"status per {REG_VERIFIED_ON}"
        rows = [[ws.cell(r, c).value for c in range(1, 7)] for r in range(5, 5 + len(REG_MAPPING))]
        assert rows == [list(x) for x in REG_MAPPING]
        assert not any("Belum diverifikasi" in str(r[2]) for r in rows)
    refs = " ".join(r[1] for r in REG_MAPPING)
    for reg in ("POJK 65/POJK.03/2016", "POJK 8/POJK.03/2014", "POJK 21/POJK.03/2014", "POJK 2 Tahun 2024", "POJK 20 Tahun 2025"):
        assert reg in refs
    assert all(src.startswith("https://") for src in REG_SOURCES)


def test_readme_results_table_matches_outputs():
    text = (ROOT / "README.md").read_text(encoding="utf-8")
    out = datastore.load_outputs()
    se = out["stress_enterprise_summary"].set_index("scenario")
    cap = out["capital_overlay"].set_index("scenario")
    f1 = lambda x: f"{x:.1f}".replace(".", ",")
    ent = " | ".join(f"{f1(scoring.r1(se.at[s, 'enterprise_score']))} {se.at[s, 'rating']}" for s in datastore.SCENARIOS)
    rau = " | ".join(f"{round(se.at[s, 'limit_utilization'])}%" for s in datastore.SCENARIOS)
    br = " | ".join(f"{se.at[s, 'tolerance_breaches']} / {se.at[s, 'limit_breaches']}" for s in datastore.SCENARIOS)
    kp = " | ".join(f"{f1(cap.at[s, 'kpmm'])}%" for s in datastore.SCENARIOS)
    for expected in (f"| Enterprise score | {ent} |", f"| Limit utilization | {rau} |",
                     f"| KRI skor ≥4 / skor 5 | {br} |", f"| KPMM | {kp} |"):
        assert expected in text, expected
