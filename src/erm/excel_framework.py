"""Excel framework (spec §10) → excel/ERM_Framework.xlsx.

Angka diambil dari data/output (engine) dan data/reference. Sheet 04–07 memuat formula Excel hidup
(threshold, skoring, agregasi, enterprise, capital overlay) yang direkonsiliasi ke hasil engine."""
from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import datastore

EXCEL_PATH = datastore.ROOT / "excel" / "ERM_Framework.xlsx"
RECALC_SCRIPT = Path(os.environ.get("ERM_RECALC_SCRIPT", "/mnt/skills/public/xlsx/scripts/recalc.py"))
DISCLAIMER = ("This project uses a fictional Islamic bank and fully synthetic data. All risk appetite thresholds, scenarios, "
              "sensitivities, and scoring rules are project-specific analytical assumptions created for educational and "
              "portfolio purposes. They do not represent any actual bank's internal limits, regulatory thresholds, or an "
              "official OJK risk profile rating.")

FONT = "Arial"
NAVY = "1F2A44"
BLUE = Font(name=FONT, size=10, color="0000FF")
GREEN_LINK = Font(name=FONT, size=10, color="008000")
BLACK = Font(name=FONT, size=10, color="000000")
BOLD = Font(name=FONT, size=10, bold=True)
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
SECTION_FONT = Font(name=FONT, size=11, bold=True, color=NAVY)
TITLE_FONT = Font(name=FONT, size=14, bold=True, color=NAVY)
NOTE_FONT = Font(name=FONT, size=9, italic=True, color="555555")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
ZONE_FILL = {"Green": "C8E6C9", "Amber": "FFF3C4", "Red": "F8CACA"}
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

S01, S02, S03, S04, S05 = "01_Risk_Taxonomy", "02_Risk_Appetite_Statement", "03_KRI_Definitions", "04_Thresholds", "05_Scoring_Methodology"
S06, S07, S08, S09, S10 = "06_Risk_Assessment", "07_Stress_Scenarios", "08_Management_Actions", "09_Regulatory_Mapping", "10_Change_Log"

# ---------------------------------------------------------------- konten kualitatif (framework)
RISK_DRIVERS = {
    "Credit": "Perlambatan ekonomi, pelonggaran underwriting UMKM, konsentrasi obligor & sektor",
    "Market": "Pergerakan nilai tukar (PDN), perubahan imbal hasil sukuk FVOCI, repricing gap",
    "Liquidity": "Penarikan DPK, konsentrasi deposan besar, maturity mismatch jangka pendek",
    "Operational": "Process failure pada transaksi digital, kegagalan sistem, fraud eksternal",
    "Strategic": "Deviasi growth vs RKB, pencapaian laba, efisiensi biaya, eksekusi inisiatif",
    "Compliance": "Pelanggaran ketentuan, remediasi terlambat, kesesuaian akad dengan opini DPS",
    "Legal": "Sengketa pembiayaan & agunan, gugatan nasabah, aging kasus",
    "Reputation": "Gangguan layanan digital, komplain severe, sentimen negatif media/sosmed",
    "Rate of Return": "Gap bagi hasil vs benchmark pasar, ekspektasi deposan, kompresi margin tetap",
    "Investment": "Konsentrasi penerbit, penurunan rating, penurunan nilai pasar, durasi portofolio",
}
APPETITE_STATEMENT = {
    "Credit": "ASB menjaga kualitas pembiayaan dengan NPF gross di bawah 2,5% dan tidak mentoleransi NPF di atas 3,5%; pertumbuhan UMKM tidak boleh dicapai dengan melonggarkan standar underwriting.",
    "Market": "ASB memiliki appetite rendah terhadap risiko pasar; posisi valas dan sensitivitas sukuk FVOCI dijaga agar dampaknya terhadap modal tetap terbatas.",
    "Liquidity": "ASB mempertahankan buffer aset likuid yang memadai untuk menghadapi penarikan DPK 30 hari dan menghindari ketergantungan pada deposan besar.",
    "Operational": "ASB tidak mentoleransi insiden high-severity berulang; kerugian operasional dijaga di bawah 1% gross income dan layanan digital harus andal.",
    "Strategic": "ASB tumbuh sesuai RKB; deviasi growth yang material, baik di atas maupun di bawah target, diperlakukan sebagai sinyal risiko.",
    "Compliance": "ASB memiliki zero appetite terhadap pelanggaran prinsip syariah; temuan DPS dan pelanggaran high-severity wajib diremediasi tepat waktu.",
    "Legal": "ASB meminimalkan eksposur hukum melalui dokumentasi akad yang kuat dan penyelesaian sengketa yang tepat waktu.",
    "Reputation": "ASB menjaga kepercayaan nasabah; komplain severe dan sentimen negatif dikelola sebelum berkembang menjadi isu publik.",
    "Rate of Return": "ASB mengelola displaced commercial risk agar imbal hasil deposan tetap kompetitif tanpa mengorbankan keberlanjutan margin.",
    "Investment": "ASB berinvestasi pada sukuk berkualitas tinggi dengan konsentrasi penerbit dan durasi yang terkendali.",
}
KRI_FORMULA = {
    "CR01": ("Σ outstanding kolektibilitas 3–5 / Σ outstanding × 100", "financing_portfolio"),
    "CR02": ("Σ outstanding UMKM kolektibilitas 3–5 / Σ outstanding UMKM × 100", "financing_portfolio"),
    "CR03": ("Σ outstanding kolektibilitas 2 / Σ outstanding × 100", "financing_portfolio"),
    "CR04": ("Σ outstanding 10 obligor terbesar / Σ outstanding × 100", "financing_portfolio"),
    "CR05": ("Σ outstanding sektor terbesar / Σ outstanding × 100", "financing_portfolio"),
    "MR01": ("Σ |net position valas| / modal KPMM × 100", "market_positions, capital_monthly"),
    "MR02": ("Σ(MV FVOCI × modified duration) × 1% / modal × 100", "investment_portfolio, capital_monthly"),
    "MR03": ("|Σ(RSA − RSL) bucket ≤12 bulan| / total aset × 100", "maturity_profile, balance_sheet_monthly"),
    "LQ01": ("(excess GWM + penempatan + sukuk unencumbered × (1 − 5%)) / net outflow 30 hari × 100; net outflow = gross outflow × (8.750 / gross outflow Sep-2026)", "balance_sheet_monthly, investment_portfolio, deposits_monthly"),
    "LQ02": ("Pembiayaan bruto / total DPK × 100", "balance_sheet_monthly"),
    "LQ03": ("Σ saldo 20 deposan terbesar / total DPK × 100", "top_depositors, balance_sheet_monthly"),
    "LQ04": ("max(0, liabilitas − aset jatuh tempo ≤1 bulan) / total aset × 100", "maturity_profile, balance_sheet_monthly"),
    "LQ05": ("max(0, (DPK t−3 − DPK t) / DPK t−3 × 100)", "balance_sheet_monthly"),
    "OR01": ("Jumlah insiden pada bulan berjalan", "operational_events"),
    "OR02": ("Jumlah insiden severity High, rolling 3 bulan", "operational_events"),
    "OR03": ("Σ net loss rolling 12 bulan / gross income TTM × 100", "operational_events, capital_monthly"),
    "OR04": ("Insiden repeat / total insiden, rolling 3 bulan × 100", "operational_events"),
    "OR05": ("Insiden melewati SLA (selesai > SLA, atau masih open dengan umur > SLA) / total, rolling 3 bulan × 100", "operational_events"),
    "ST01": ("|growth pembiayaan YoY − target 15%|", "strategic_monthly"),
    "ST02": ("Laba YTD / anggaran laba YTD (RKB) × 100", "capital_monthly"),
    "ST03": ("Cost to income (proxy BOPO)", "strategic_monthly"),
    "ST04": ("Inisiatif terlambat / total inisiatif × 100", "strategic_monthly"),
    "CP01": ("Jumlah pelanggaran sumber regulator + internal audit, rolling 3 bulan", "compliance_events"),
    "CP02": ("Pelanggaran High yang terbuka dan melewati target remediasi pada akhir bulan", "compliance_events"),
    "CP03": ("Event repeat / total event kepatuhan, rolling 12 bulan × 100", "compliance_events"),
    "CP04": ("Temuan DPS terbuka dengan umur > 90 hari (dihitung dari awal bulan temuan)", "compliance_events"),
    "LG01": ("Jumlah kasus terbuka pada akhir bulan", "legal_cases"),
    "LG02": ("Jumlah kasus terbuka severity High", "legal_cases"),
    "LG03": ("Σ potensi kerugian kasus terbuka / modal × 100", "legal_cases, capital_monthly"),
    "LG04": ("Kasus terbuka berumur > 365 hari / kasus terbuka × 100", "legal_cases"),
    "RP01": ("(Jumlah komplain bulan t / bulan t−1 − 1) × 100", "complaints"),
    "RP02": ("Komplain severity Severe / total komplain × 100", "complaints"),
    "RP03": ("Komplain selesai melewati SLA / total komplain × 100", "complaints"),
    "RP04": ("Σ negative mentions / Σ mentions × 100", "sentiment_monthly"),
    "RR01": ("Benchmark pasar deposito − equivalent rate deposito mudharabah ASB (pp)", "deposits_monthly"),
    "RR02": ("Return ekspektasi deposan − realisasi (pp)", "rate_of_return_monthly"),
    "RR03": ("NIM syariah", "rate_of_return_monthly"),
    "RR04": ("Σ outstanding margin tetap dengan sisa tenor > 36 bulan / Σ outstanding × 100", "financing_portfolio"),
    "IV01": ("MV penerbit non-sovereign terbesar / MV portofolio × 100", "investment_portfolio"),
    "IV02": ("MV rating di bawah idA / MV portofolio × 100", "investment_portfolio"),
    "IV03": ("max(0, (1 − MV FVOCI t / MV FVOCI t−3) × 100)", "investment_portfolio"),
    "IV04": ("Σ(MV × modified duration) / Σ MV", "investment_portfolio"),
    "CAP01": ("Modal KPMM / ATMR × 100 (overlay; tidak masuk risk score)", "capital_monthly"),
}
REG_VERIFIED_ON = "2026-09-17"
REG_HEADERS = ["area", "rujukan", f"status per {REG_VERIFIED_ON}", "dasar verifikasi", "dipakai di", "catatan"]
# Hasil verifikasi status peraturan (penelusuran JDIH BPK & laman regulasi OJK). Status dapat berubah;
# konfirmasi ulang ke JDIH OJK sebelum dikutip.
REG_MAPPING = [
    ("Manajemen risiko BUS/UUS (10 jenis risiko)", "POJK 65/POJK.03/2016",
     "Berlaku",
     "JDIH BPK: status Berlaku, efektif 28-12-2016; mencabut PBI 13/23/PBI/2011; tidak tercatat perubahan/pencabutan",
     "01_Risk_Taxonomy, 02, 04", "Area inspirasi taksonomi 10 risiko; threshold project bukan ketentuan"),
    ("Penilaian tingkat kesehatan BUS/UUS (RBBR)", "POJK 8/POJK.03/2014 dan SEOJK 10/SEOJK.03/2014",
     "Berlaku (POJK); SEOJK masih tercantum di laman OJK",
     "JDIH BPK: POJK 8/2014 status Berlaku, efektif 13-06-2014, tidak tercatat pengganti; SEOJK 10/2014 tercantum di laman regulasi OJK, status pencabutan tidak ditemukan",
     "05_Scoring_Methodology (label rating)", "Bukan replika metodologi penilaian tingkat kesehatan bank"),
    ("KPMM BUS", "POJK 21/POJK.03/2014",
     "Berlaku; perubahan tidak ditemukan",
     "JDIH BPK: status Berlaku, efektif 01-01-2015; tidak tercatat perubahan. Terkait: POJK 21 Tahun 2025 (leverage ratio BUS, minimum 3%) berlaku 17-09-2025",
     "07_Stress_Scenarios (capital overlay)", "Frasa 'dan perubahannya' di spec §11 dihapus karena perubahan tidak ditemukan"),
    ("Tata kelola syariah / DPS", "POJK 2 Tahun 2024 (tata kelola syariah BUS/UUS); melengkapi POJK 17 Tahun 2023 (tata kelola bank umum)",
     "Berlaku",
     "Laman regulasi OJK: berlaku 16-02-2024; mengatur DPS, fungsi kepatuhan syariah, manajemen risiko syariah, audit intern syariah, kaji ulang syariah eksternal",
     "03 (CP04), 06", "Menggantikan status 'periksa regulasi terbaru' di spec §11"),
    ("Likuiditas (LCR/NSFR) BUS/UUS", "POJK 20 Tahun 2025",
     "Berlaku (pemenuhan bertahap)",
     "JDIH BPK & laman OJK: ditetapkan 12-09-2025, berlaku 17-09-2025; LCR minimum 80% (30-06-2026) → 90% (30-06-2027) → 100% (30-06-2028); NSFR 80% (31-12-2026) → 90% (31-12-2027) → 100% (31-12-2028); publikasi pertama Sep-2026",
     "03 (LQ01)", "SLSI project bukan LCR regulatori; per reporting date Sep-2026 minimum LCR regulatori 80%"),
    ("Posisi Devisa Neto (PDN)", "PBI 5/13/PBI/2003 sebagaimana diubah terakhir dengan PBI 12/10/PBI/2010",
     "Belum terkonfirmasi",
     "Tercantum di laman regulasi OJK; tidak ditemukan POJK pengganti, tetapi penerapan pada BUS dan status terkini belum dapat dipastikan dari sumber publik yang ditelusuri",
     "03 (MR01)", "Wajib dicek ke JDIH OJK sebelum dikutip"),
]
REG_SOURCES = [
    "https://peraturan.bpk.go.id/Details/128410/peraturan-ojk-no-65pojk032016-tahun-2016",
    "https://peraturan.bpk.go.id/Details/129940/peraturan-ojk-no-8pojk032014-tahun-2014",
    "https://ojk.go.id/id/regulasi/Pages/SEOJK-tentang-Penilaian-Tingkat-Kesehatan-Bank-Umum-Syariah-dan-Unit-Usaha-Syariah.aspx",
    "https://peraturan.bpk.go.id/Details/129868/peraturan-ojk-no-21pojk032014-tahun-2014",
    "https://ojk.go.id/id/regulasi/Pages/POJK-2-Tahun-2024-Penerapan-Tata-Kelola-Syariah-Bagi-Bank-Umum-Syariah-dan-Unit-Usaha-Syariah.aspx",
    "https://peraturan.bpk.go.id/Details/330102/peraturan-ojk-no-20-tahun-2025",
    "https://ojk.go.id/id/regulasi/Pages/POJK-20-Tahun-2025-Kewajiban-Pemenuhan-Rasio-Kecukupan-Likuiditas-dan-Rasio-Pendanaan-Stabil-Bersih-Bagi-BUS-UUS.aspx",
    "https://www.antaranews.com/berita/5210685/ojk-rilis-pojk-bagi-bank-syariah-guna-perkuat-likuiditas-permodalan",
    "https://ojk.go.id/id/regulasi/Pages/PBI-tentang-Perubahan-Ketiga-atas-PBI-Nomor-513PBI2003-tentang-Posisi-Devisa-Neto-Bank-Umum.aspx",
]


def _write_reg_table(ws, start_row: int = 4) -> int:
    _header(ws, start_row, REG_HEADERS)
    for i, row in enumerate(REG_MAPPING):
        r = start_row + 1 + i
        for j, v in enumerate(row):
            fill = None
            if j == 2:
                fill = PatternFill("solid", fgColor=ZONE_FILL["Green" if v.startswith("Berlaku") else "Amber"])
            _put(ws, r, j + 1, v, wrap=True, fill=fill)
        ws.row_dimensions[r].height = 62
    r = start_row + len(REG_MAPPING) + 2
    ws.cell(row=r, column=1, value="Sumber verifikasi").font = SECTION_FONT
    for k, url in enumerate(REG_SOURCES):
        ws.cell(row=r + 1 + k, column=1, value=url).font = NOTE_FONT
    ws.cell(row=r + 2 + len(REG_SOURCES), column=1,
            value="Status dapat berubah setelah tanggal verifikasi; konfirmasi ulang ke JDIH OJK sebelum dikutip. "
                  "Pemetaan hanya menjelaskan area yang menginspirasi; tidak ada klaim threshold project sama dengan ketentuan.").font = NOTE_FONT
    _widths(ws, {"A": 30, "B": 38, "C": 22, "D": 60, "E": 26, "F": 40})
    return r + 2 + len(REG_SOURCES)
CHANGE_LOG = [
    ("v3.0", "2026-01-15", "Threshold, bobot KRI, bobot materialitas, dan aturan skoring ditetapkan (ERM Blueprint v3)", "Risk Management Committee (fiktif)", "Spec §3–§4"),
    ("v3.0", "2026-01-15", "Skenario makro dan koefisien transmisi stress ditetapkan", "Risk Management Committee (fiktif)", "Spec §5"),
    ("v3.1", "2026-09-16", "Katalog final 42 KRI + 1 capital overlay (CAP01)", "Project owner (fiktif)", "Keputusan Fase 2"),
    ("v3.1", "2026-09-16", "SLSI: net outflow 30 hari 8.750 dipakai tetap sebagai anchor Sep-2026", "Project owner (fiktif)", "Keputusan Fase 2"),
    ("v3.2", "2026-09-17", "Aturan eskalasi enterprise (skor ≥ 4,5 & bobot ≥ 0,10) memakai presisi penuh", "Project owner (fiktif)", "Keputusan Fase 3"),
    ("v3.2", "2026-09-17", "Nilai KRI dibulatkan 6 desimal (presisi pengukuran) sebelum skoring", "Project owner (fiktif)", "Keputusan Fase 4"),
    ("v3.3", "2026-09-17", "Regulatory mapping diverifikasi: POJK 65/2016, 8/2014, 21/2014 berlaku; tata kelola syariah = POJK 2/2024; LCR/NSFR BUS = POJK 20/2025; PDN belum terkonfirmasi", "Project owner (fiktif)", "Sheet 09"),
    ("v3.2", "2026-09-30", "Judgment override Compliance +0,3 (model 3,3 → final 3,6), berlaku s.d. 2026-12", "Risk Management Committee (fiktif)", "judgment_overrides.csv"),
]


# ---------------------------------------------------------------- helper
def _title(ws, title: str, subtitle: str | None = None):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle or "PT Amanah Syariah Bank (fiktif) · reporting date 30-Sep-2026 · data sintetis"
    ws["A2"].font = NOTE_FONT


def _section(ws, row: int, text: str):
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT


def _header(ws, row: int, headers: list[str], col: int = 1):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _put(ws, row: int, col: int, value, font=BLACK, fmt: str | None = None, fill=None, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font, c.border = font, BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    return c


def _widths(ws, widths: dict[str, float]):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _q(sheet: str) -> str:
    return f"'{sheet}'"


def _rating_formula(cell: str) -> str:
    return (f'IF({cell}<1.5,"Low",IF({cell}<2.5,"Low to Moderate",IF({cell}<3.5,"Moderate",'
            f'IF({cell}<4.5,"Moderate to High","High"))))')


def _score_formula(x: str, d: str, c1: str, c2: str, c3: str, c4: str) -> str:
    ax = f"ABS({x})"
    return (f'IF({d}="lower",IF({ax}<={c1},1,IF({ax}<={c2},2,IF({ax}<={c3},3,IF({ax}<={c4},4,5)))),'
            f'IF({ax}>={c1},1,IF({ax}>={c2},2,IF({ax}>={c3},3,IF({ax}>={c4},4,5)))))')


def describe_rule(kri: str, rule: dict) -> tuple[str, str]:
    r = rule["rule"]
    if r == "additive":
        parts = []
        if "gdp_coef" in rule:
            parts.append(f"{rule['gdp_coef']:+g} × ΔGDP")
        if "bps_coef" in rule:
            parts.append(f"{rule['bps_coef']:+g} × Δbps/100")
        return "Nilai " + " ".join(parts), ", ".join(f"{k}={v}" for k, v in rule.items() if k != "rule")
    text = {
        "multiple_of_delta": f"Nilai + {rule.get('multiplier')} × Δ{rule.get('ref_kri')}",
        "scale_by_variable": f"Nilai × (1 + {rule.get('variable')})",
        "scale_by_variable_ceil": f"ceil(Nilai × (1 + {rule.get('variable')}))",
        "scale_bps": "Nilai × (100 + Δbps)/100",
        "slsi_recalc": "Hitung ulang: haircut = 5% + |MV shock|; net outflow + run-off tambahan × DPK",
        "fdr_recalc": "Pembiayaan / (DPK × (1 − penurunan DPK 3 bln))",
        "divide_by_dpk_decline": "Nilai ÷ (1 − penurunan DPK 3 bln)",
        "additive_variable": f"Nilai + {rule.get('coef')} × {rule.get('variable')}",
        "max_baseline_variable": f"max(Nilai, {rule.get('variable')})",
        "max_baseline_abs_variable": f"max(Nilai, |{rule.get('variable')}|)",
        "growth_deviation": f"|growth aktual − {rule.get('gdp_coef')} × ΔGDP − {rule.get('growth_target')}|",
    }[r]
    return text, ", ".join(f"{k}={v}" for k, v in rule.items() if k != "rule")


# ---------------------------------------------------------------- builder
def build_workbook(outputs: dict | None = None, ref: dict | None = None, cfg: dict | None = None) -> Workbook:
    outputs = outputs or datastore.load_outputs()
    ref = ref or datastore.load_reference()
    cfg = cfg or datastore.load_config()
    rep = cfg["bank_profile"]["reporting"]["period_end"]
    ra, rw = ref["risk_appetite"], ref["risk_weights"]
    risks = list(rw.risk_type)
    n_kri = len(ra)
    km_sep = outputs["kri_monthly"].query("period == @rep").set_index("kri_id")
    rs_sep = outputs["risk_scores_monthly"].query("period == @rep").set_index("risk_type")
    ep_sep = outputs["enterprise_profile_monthly"].query("period == @rep").iloc[0]

    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in (S01, S02, S03, S04, S05, S06, S07, S08, S09, S10)}

    # ---- 04_Thresholds (dibangun lebih dulu karena dirujuk sheet lain)
    ws = sheets[S04]
    _title(ws, "04 · Thresholds (risk_appetite.csv)", "Input terkontrol: ubah cut-point/bobot hanya melalui change log (sheet 10). Sel biru = input.")
    cols = list(ra.columns)
    _header(ws, 4, cols + ["cek_urutan_cutpoint", "Σ_bobot_risk"])
    T0, T1 = 5, 5 + n_kri - 1
    for i, row in enumerate(ra.itertuples(index=False)):
        r = T0 + i
        for j, v in enumerate(row):
            v = v.item() if hasattr(v, "item") else v
            _put(ws, r, j + 1, v, font=BLUE if cols[j] not in ("kri_id", "risk_type", "kri_name") else BLACK)
        _put(ws, r, len(cols) + 1, f'=IF(E{r}="lower",AND(F{r}<=G{r},G{r}<=H{r},H{r}<=I{r}),AND(F{r}>=G{r},G{r}>=H{r},H{r}>=I{r}))')
        _put(ws, r, len(cols) + 2, f"=SUMIF($B${T0}:$B${T1},B{r},$J${T0}:$J${T1})", fmt="0.00")
    dv_dir = DataValidation(type="list", formula1='"lower,higher"', allow_blank=False)
    dv_w = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
    dv_s = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    dv_yn = DataValidation(type="list", formula1='"Y,N"')
    dv_m = DataValidation(type="list", formula1='"quantitative,semi-quantitative,KRI-based"')
    for dv, colname in ((dv_dir, "direction"), (dv_w, "kri_weight"), (dv_s, "target_score"), (dv_yn, "stress_sensitive"),
                        (dv_m, "measurement_type")):
        letter = get_column_letter(cols.index(colname) + 1)
        dv.add(f"{letter}{T0}:{letter}{T1}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "B5"
    _widths(ws, {"A": 8, "B": 14, "C": 44, "D": 12, "E": 9, "Q": 26, "R": 14, "S": 12})
    TH = _q(S04)

    def th(col_letter: str, key_cell: str) -> str:
        return f"INDEX({TH}!${col_letter}${T0}:${col_letter}${T1},MATCH({key_cell},{TH}!$A${T0}:$A${T1},0))"

    # ---- 06_Risk_Assessment
    ws = sheets[S06]
    _title(ws, f"06 · Risk Assessment — snapshot {rep} (hasil engine + formula hidup)",
           "Nilai KRI, skor engine, trend, dan early warning dari data/output. Kolom formula menghitung ulang skor; kolom override (kuning) dapat diubah.")
    _section(ws, 4, "A. KRI snapshot")
    _header(ws, 5, ["kri_id", "risk_type", "kri_name", "unit", "direction", "c1", "c2", "c3", "c4", "kri_weight",
                    "value (engine)", "score (engine)", "score (formula)", "zone (formula)", "utilization (formula)",
                    "trend_label (engine)", "early_warning (engine)", "cek skor"])
    A0, A1 = 6, 6 + n_kri - 1
    for i, kid in enumerate(ra.kri_id):
        r = A0 + i
        k = km_sep.loc[kid]
        _put(ws, r, 1, kid)
        _put(ws, r, 2, ra.set_index("kri_id").at[kid, "risk_type"])
        _put(ws, r, 3, ra.set_index("kri_id").at[kid, "kri_name"], wrap=False)
        _put(ws, r, 4, ra.set_index("kri_id").at[kid, "unit"])
        for col, src in zip("EFGHIJ", "EFGHIJ"):
            _put(ws, r, " ABCDEFGHIJ".index(col), "=" + th(src, f"$A{r}"), font=GREEN_LINK,
                 fmt="0.00" if col != "E" else None)
        _put(ws, r, 11, float(k.value), font=BLUE, fmt="0.0000")
        _put(ws, r, 12, int(k.score), font=BLUE)
        _put(ws, r, 13, "=" + _score_formula(f"K{r}", f"E{r}", f"F{r}", f"G{r}", f"H{r}", f"I{r}"))
        _put(ws, r, 14, f'=CHOOSE(M{r},"Strong","Within appetite","Watch","Tolerance breach","Limit breach")')
        _put(ws, r, 15, f'=MIN(1.5,IF(E{r}="lower",ABS(K{r})/I{r},I{r}/K{r}))', fmt="0.0%")
        _put(ws, r, 16, k.trend_label, font=BLUE)
        _put(ws, r, 17, bool(k.early_warning_flag), font=BLUE)
        _put(ws, r, 18, f"=M{r}-L{r}")
        ws.cell(row=r, column=12).fill = PatternFill("solid", fgColor=ZONE_FILL[k.color])

    B_HDR = A1 + 3
    _section(ws, B_HDR - 1, "B. Risk score per risk type (§3.2–3.4) — formula hidup")
    _header(ws, B_HDR, ["risk_type", "materiality_weight", "weighted (formula)", "max KRI score", "floor = max − 1,5",
                        "model = max(weighted, floor)", "override Δ (input)", "final = model + override (1–5)",
                        "final (tampilan, ROUND 1)", "rating (formula)", "final (engine)", "selisih formula − engine",
                        "RAU per risk (formula)"])
    B0, B1 = B_HDR + 1, B_HDR + len(risks)
    dv_ov = DataValidation(type="decimal", operator="between", formula1="-1", formula2="1")
    ws.add_data_validation(dv_ov)
    for i, risk in enumerate(risks):
        r = B0 + i
        rsr = rs_sep.loc[risk]
        _put(ws, r, 1, risk)
        _put(ws, r, 2, f"=INDEX({_q(S01)}!$F$6:$F$15,MATCH(A{r},{_q(S01)}!$A$6:$A$15,0))", font=GREEN_LINK, fmt="0.00")
        _put(ws, r, 3, f"=SUMPRODUCT(($B${A0}:$B${A1}=A{r})*$M${A0}:$M${A1}*$J${A0}:$J${A1})", fmt="0.0000")
        _put(ws, r, 4, f"=SUMPRODUCT(MAX(($B${A0}:$B${A1}=A{r})*$M${A0}:$M${A1}))")
        _put(ws, r, 5, f"=D{r}-1.5", fmt="0.0000")
        _put(ws, r, 6, f"=MAX(C{r},E{r})", fmt="0.0000")
        c = _put(ws, r, 7, float(rsr.override_delta), font=BLUE, fmt="+0.0;-0.0;0.0", fill=INPUT_FILL)
        dv_ov.add(c.coordinate)
        _put(ws, r, 8, f"=MIN(5,MAX(1,F{r}+G{r}))", fmt="0.0000")
        _put(ws, r, 9, f"=ROUND(H{r},1)", fmt="0.0")
        _put(ws, r, 10, "=" + _rating_formula(f"I{r}"))
        _put(ws, r, 11, float(rsr.final_score), font=BLUE, fmt="0.0000")
        _put(ws, r, 12, f"=H{r}-K{r}", fmt="0.0000")
        _put(ws, r, 13, f"=SUMPRODUCT(($B${A0}:$B${A1}=A{r})*$O${A0}:$O${A1}*$J${A0}:$J${A1})", fmt="0.0000")
    ws.cell(row=B0 + risks.index("Compliance"), column=7).comment = Comment(
        "Override judgment dari judgment_overrides.csv (RMC fiktif, berlaku s.d. 2026-12). Batas ±1,0.", "ERM engine")

    C_HDR = B1 + 3
    _section(ws, C_HDR - 1, "C. Enterprise risk score (§3.5–3.6) — formula hidup")
    _header(ws, C_HDR, ["Komponen", "Formula", "Engine", "Selisih"])
    C0 = C_HDR + 1
    r0, r1, r2, r3, r4 = C0, C0 + 1, C0 + 2, C0 + 3, C0 + 4
    er_start = 12 + max(len(outputs["management_actions"]), 1) + 4
    rows_c = [
        ("Σ final × bobot materialitas", f"=SUMPRODUCT(H{B0}:H{B1},B{B0}:B{B1})", None, "0.0000"),
        ("Floor enterprise = max(final) − 1,5", f"=MAX(H{B0}:H{B1})-1.5", None, "0.0000"),
        ("Setelah floor", f"=MAX(B{r0},B{r1})", None, "0.0000"),
        ("Eskalasi: ada risk bobot ≥ 0,10 dan skor ≥ 4,5 → minimal 3,5 (presisi penuh) = ENTERPRISE SCORE",
         f'=IF(COUNTIFS(H{B0}:H{B1},">=4.5",B{B0}:B{B1},">=0.1")>0,MAX(B{r2},3.5),B{r2})', float(ep_sep.enterprise_score), "0.0000"),
        ("Enterprise score (tampilan, ROUND 1)", f"=ROUND(B{r3},1)", None, "0.0"),
        ("Enterprise rating", "=" + _rating_formula(f"B{r4}"), ep_sep.rating, None),
        ("Limit utilization (%)", f"=SUMPRODUCT(M{B0}:M{B1},B{B0}:B{B1})*100", float(ep_sep.limit_utilization), "0.0000"),
        ("Tolerance breach (KRI skor ≥ 4)", f'=COUNTIFS($M${A0}:$M${A1},">=4",$B${A0}:$B${A1},"<>Capital Overlay")', int(ep_sep.tolerance_breaches), "0"),
        ("Limit breach (KRI skor = 5)", f'=COUNTIFS($M${A0}:$M${A1},5,$B${A0}:$B${A1},"<>Capital Overlay")', int(ep_sep.limit_breaches), "0"),
        ("Emerging risk (register)", f"=COUNTA({_q(S08)}!$A${er_start + 2}:$A${er_start + 1 + len(ref['emerging_risks'])})", int(ep_sep.emerging_count), "0"),
    ]
    for i, (label, formula, engine, fmt) in enumerate(rows_c):
        r = C0 + i
        _put(ws, r, 1, label, font=BOLD if i in (3, 4, 5) else BLACK, wrap=True)
        _put(ws, r, 2, formula, fmt=fmt)
        if engine is not None:
            _put(ws, r, 3, engine, font=BLUE, fmt=fmt)
            _put(ws, r, 4, f'=IF(ISNUMBER(C{r}),B{r}-C{r},IF(B{r}=C{r},0,1))', fmt="0.0000")
    refs = {"r3": r3, "r4": r4}
    ws.freeze_panes = "B6"
    _widths(ws, {"A": 34, "B": 16, "C": 40, "D": 12, "E": 10, "F": 10, "G": 12, "H": 14, "I": 12, "J": 16, "K": 14,
                 "L": 14, "M": 14, "N": 16, "O": 14, "P": 18, "Q": 14, "R": 9})
    RA_ = _q(S06)
    ENT_ROW = refs["r3"]
    ENT_DISPLAY_ROW = refs["r4"]

    # ---- 01_Risk_Taxonomy
    ws = sheets[S01]
    _title(ws, "01 · Risk Taxonomy — 10 risk type Bank Umum Syariah")
    ws["A3"] = DISCLAIMER
    ws["A3"].font = NOTE_FONT
    ws["A3"].alignment = WRAP
    ws.merge_cells("A3:H3")
    ws.row_dimensions[3].height = 48
    _header(ws, 5, ["risk_type", "definisi", "driver utama", "measurement approach", "depth", "materiality weight",
                    "risk owner (KRI owner dominan)", "jumlah KRI (formula)"])
    owners = ra[ra.risk_type != "Capital Overlay"].groupby("risk_type").kri_owner.agg(lambda s: s.value_counts().index[0])
    from .export_powerbi import RISK_DEFINITIONS
    for i, row in enumerate(rw.itertuples(index=False)):
        r = 6 + i
        _put(ws, r, 1, row.risk_type, font=BOLD)
        _put(ws, r, 2, RISK_DEFINITIONS[row.risk_type], wrap=True)
        _put(ws, r, 3, RISK_DRIVERS[row.risk_type], wrap=True)
        _put(ws, r, 4, row.measurement_approach)
        _put(ws, r, 5, row.depth)
        _put(ws, r, 6, float(row.materiality_weight), font=BLUE, fmt="0.00")
        _put(ws, r, 7, owners[row.risk_type], wrap=True)
        _put(ws, r, 8, f"=COUNTIF({TH}!$B${T0}:$B${T1},A{r})")
    _put(ws, 16, 5, "Σ bobot", font=BOLD)
    _put(ws, 16, 6, "=SUM(F6:F15)", fmt="0.00")
    _put(ws, 16, 8, "=SUM(H6:H15)")
    _widths(ws, {"A": 16, "B": 48, "C": 48, "D": 20, "E": 13, "F": 12, "G": 26, "H": 12})

    # ---- 02_Risk_Appetite_Statement
    ws = sheets[S02]
    _title(ws, "02 · Risk Appetite Statement")
    _header(ws, 4, ["risk_type", "pernyataan appetite (kualitatif)", "metrik utama (KRI bobot terbesar)", "nama KRI",
                    "appetite (c2)", "early warning trigger (c3)", "limit (c4)", f"nilai {rep}", f"skor KRI {rep}",
                    f"final risk score {rep}"])
    main_kri = ra[ra.risk_type != "Capital Overlay"].sort_values(["risk_type", "kri_weight"], ascending=[True, False]) \
        .groupby("risk_type").kri_id.first()
    for i, risk in enumerate(risks):
        r = 5 + i
        _put(ws, r, 1, risk, font=BOLD)
        _put(ws, r, 2, APPETITE_STATEMENT[risk], wrap=True)
        _put(ws, r, 3, main_kri[risk])
        _put(ws, r, 4, "=" + th("C", f"$C{r}"), font=GREEN_LINK, wrap=True)
        _put(ws, r, 5, "=" + th("G", f"$C{r}"), font=GREEN_LINK, fmt="0.00")
        _put(ws, r, 6, "=" + th("H", f"$C{r}"), font=GREEN_LINK, fmt="0.00")
        _put(ws, r, 7, "=" + th("I", f"$C{r}"), font=GREEN_LINK, fmt="0.00")
        _put(ws, r, 8, f"=INDEX({RA_}!$K${A0}:$K${A1},MATCH($C{r},{RA_}!$A${A0}:$A${A1},0))", font=GREEN_LINK, fmt="0.00")
        _put(ws, r, 9, f"=INDEX({RA_}!$M${A0}:$M${A1},MATCH($C{r},{RA_}!$A${A0}:$A${A1},0))", font=GREEN_LINK)
        _put(ws, r, 10, f"=INDEX({RA_}!$I${B0}:$I${B1},MATCH($A{r},{RA_}!$A${B0}:$A${B1},0))", font=GREEN_LINK, fmt="0.0")
        ws.row_dimensions[r].height = 48
    _widths(ws, {"A": 16, "B": 60, "C": 12, "D": 34, "E": 12, "F": 14, "G": 10, "H": 12, "I": 10, "J": 12})

    # ---- 03_KRI_Definitions
    ws = sheets[S03]
    _title(ws, "03 · KRI Definitions — formula, sumber data, frekuensi, owner")
    _header(ws, 4, ["kri_id", "risk_type", "kri_name", "unit", "arah", "formula operasional", "sumber data (data/raw)",
                    "frekuensi", "measurement type", "owner", "stress sensitive"])
    for i, row in enumerate(ra.itertuples(index=False)):
        r = 5 + i
        formula, source = KRI_FORMULA[row.kri_id]
        for j, v in enumerate([row.kri_id, row.risk_type, row.kri_name, row.unit, row.direction, formula, source,
                               row.frequency, row.measurement_type, row.kri_owner, row.stress_sensitive]):
            _put(ws, r, j + 1, v, wrap=j in (2, 5, 6))
    _widths(ws, {"A": 8, "B": 14, "C": 36, "D": 12, "E": 8, "F": 60, "G": 34, "H": 10, "I": 16, "J": 24, "K": 9})
    ws.freeze_panes = "B5"

    # ---- 05_Scoring_Methodology
    ws = sheets[S05]
    _title(ws, "05 · Scoring Methodology (§3) + contoh hitung Credit dengan formula hidup")
    _section(ws, 4, "Skala skor KRI (§3.1)")
    _header(ws, 5, ["skor", "zona", "makna governance", "warna"])
    for i, (s, z, m, col) in enumerate([
        (1, "Strong", "Jauh di dalam appetite (≤ c1 / ≥ c1)", "Green"), (2, "Within appetite", "c2 = Risk Appetite", "Green"),
        (3, "Watch", "Di atas appetite, di bawah c3 = Early Warning Trigger", "Amber"),
        (4, "Tolerance breach", "Melewati trigger, di bawah c4 = Risk Limit", "Red"),
        (5, "Limit breach", "Melewati limit → eskalasi ke Direksi/RMC", "Red")]):
        for j, v in enumerate([s, z, m, col]):
            _put(ws, 6 + i, j + 1, v, fill=PatternFill("solid", fgColor=ZONE_FILL[col]) if j == 3 else None)
    _section(ws, 12, "Rating (§3.4) — diterapkan pada skor dibulatkan 1 desimal ROUND_HALF_UP")
    _header(ws, 13, ["batas bawah", "batas atas (<)", "rating", "warna"])
    for i, (lo, hi, lab, col) in enumerate([(1.0, 1.5, "Low", "Green"), (1.5, 2.5, "Low to Moderate", "Green"),
                                            (2.5, 3.5, "Moderate", "Amber"), (3.5, 4.5, "Moderate to High", "Red"),
                                            (4.5, 5.0, "High", "Red")]):
        for j, v in enumerate([lo, hi, lab, col]):
            _put(ws, 14 + i, j + 1, v, fmt="0.0" if j < 2 else None)
    _section(ws, 20, "Aturan agregasi")
    rules_txt = [
        "KRI score: lower → x ≤ c1:1; ≤ c2:2; ≤ c3:3; ≤ c4:4; selain itu 5. higher → x ≥ c1:1 … KRI bertanda memakai nilai absolut.",
        "Risk: weighted = Σ(skor × bobot KRI); floor = max(skor) − 1,5; model = max(weighted, floor); final = model + override (dibatasi 1–5).",
        "Override: maksimal ±1,0, tercatat di judgment_overrides.csv, delta dibawa ke semua skenario stress.",
        "Enterprise: Σ(final × bobot materialitas); floor = max(final) − 1,5; bila ada risk bobot ≥ 0,10 dengan skor ≥ 4,5 (presisi penuh) → minimal 3,5.",
        "Utilization per KRI: |x|/c4 (lower) atau c4/x (higher), maks 150%; diagregasi dengan bobot KRI lalu bobot materialitas.",
        "Trend: rata-rata 3 bulan terakhir − 3 bulan sebelumnya, dibagi |c4 − c1|; > +10% memburuk = Deteriorating; < −10% = Improving.",
        "Early warning: skor ≤ 3 dan proyeksi garis OLS 6 bulan melewati cut-point berikutnya dalam ≤ 3 bulan.",
        "Presisi: seluruh perhitungan presisi penuh; pembulatan 1 desimal hanya untuk rating dan tampilan.",
    ]
    for i, t in enumerate(rules_txt):
        _put(ws, 21 + i, 1, f"{i + 1}.")
        ws.merge_cells(start_row=21 + i, start_column=2, end_row=21 + i, end_column=10)
        _put(ws, 21 + i, 2, t, wrap=True)
        ws.row_dimensions[21 + i].height = 28
    EX = 31
    _section(ws, EX - 1, f"Contoh hitung: Credit risk {rep} (semua sel hitung = formula; nilai KRI ditautkan dari 06)")
    _header(ws, EX, ["kri_id", "direction", "c1", "c2", "c3", "c4", "bobot", "nilai KRI", "skor (formula)",
                     "skor × bobot", "utilization", "zona"])
    credit = list(ra[ra.risk_type == "Credit"].kri_id)
    E0 = EX + 1
    for i, kid in enumerate(credit):
        r = E0 + i
        _put(ws, r, 1, kid)
        for j, src in enumerate("EFGHIJ"):
            _put(ws, r, 2 + j, "=" + th(src, f"$A{r}"), font=GREEN_LINK, fmt=None if src == "E" else "0.00")
        _put(ws, r, 8, f"=INDEX({RA_}!$K${A0}:$K${A1},MATCH($A{r},{RA_}!$A${A0}:$A${A1},0))", font=GREEN_LINK, fmt="0.0000")
        _put(ws, r, 9, "=" + _score_formula(f"H{r}", f"B{r}", f"C{r}", f"D{r}", f"E{r}", f"F{r}"))
        _put(ws, r, 10, f"=I{r}*G{r}", fmt="0.00")
        _put(ws, r, 11, f'=MIN(1.5,IF(B{r}="lower",ABS(H{r})/F{r},F{r}/H{r}))', fmt="0.0%")
        _put(ws, r, 12, f'=CHOOSE(I{r},"Strong","Within appetite","Watch","Tolerance breach","Limit breach")')
    E1 = E0 + len(credit) - 1
    steps = [
        ("Σ bobot KRI (harus 1,00)", f"=SUM(G{E0}:G{E1})", "0.00"),
        ("Weighted = Σ skor × bobot", f"=SUM(J{E0}:J{E1})", "0.0000"),
        ("Max skor KRI", f"=MAX(I{E0}:I{E1})", "0"),
        ("Floor = max − 1,5", "=B{m}-1.5", "0.0000"),
        ("Model = max(weighted, floor)", "=MAX(B{w},B{f})", "0.0000"),
        ("Override Δ (dari 06)", f'=INDEX({RA_}!$G${B0}:$G${B1},MATCH("Credit",{RA_}!$A${B0}:$A${B1},0))', "0.0"),
        ("Final = min(5, max(1, model + override))", "=MIN(5,MAX(1,B{mo}+B{ov}))", "0.0000"),
        ("Final (tampilan, ROUND 1)", "=ROUND(B{fi},1)", "0.0"),
        ("Rating", "=" + _rating_formula("B{fd}"), None),
        ("Limit utilization Credit (Σ utilization × bobot)", f"=SUMPRODUCT(K{E0}:K{E1},G{E0}:G{E1})", "0.0%"),
        ("Final Credit (engine, data/output)", float(rs_sep.loc["Credit"].final_score), "0.0000"),
        ("Selisih formula − engine", "=B{fi}-B{en}", "0.0000"),
    ]
    S0 = E1 + 2
    idx = {"w": S0 + 1, "m": S0 + 2, "f": S0 + 3, "mo": S0 + 4, "ov": S0 + 5, "fi": S0 + 6, "fd": S0 + 7, "en": S0 + 10}
    for i, (label, formula, fmt) in enumerate(steps):
        r = S0 + i
        _put(ws, r, 1, label, font=BOLD if i in (6, 7, 8) else BLACK)
        val = formula.format(**idx) if isinstance(formula, str) else formula
        _put(ws, r, 2, val, font=BLUE if not isinstance(formula, str) else BLACK, fmt=fmt)
    _widths(ws, {"A": 40, "B": 16, "C": 30, "D": 10, "E": 10, "F": 10, "G": 10, "H": 12, "I": 12, "J": 12, "K": 12, "L": 16})

    # ---- 07_Stress_Scenarios
    ws = sheets[S07]
    params, macro = cfg["stress_parameters"], ref["macro_scenarios"]
    _title(ws, "07 · Stress Scenarios (§5.1–5.5)", "Horizon 12 bulan, neraca statis, tanpa management action (gross stress impact).")
    _section(ws, 4, "A. Skenario makro (§5.1) — input")
    _header(ws, 5, ["variable", "unit", "baseline", "adverse", "severe"])
    M0 = 6
    for i, row in enumerate(macro.itertuples(index=False)):
        r = M0 + i
        _put(ws, r, 1, row.variable)
        _put(ws, r, 2, row.unit)
        for j, v in enumerate([row.baseline, row.adverse, row.severe]):
            _put(ws, r, 3 + j, float(v), font=BLUE, fmt="0.0")
    M1 = M0 + len(macro) - 1

    def mac(var: str, col: str) -> str:
        return f'INDEX({col}${M0}:{col}${M1},MATCH("{var}",$A${M0}:$A${M1},0))'

    R_HDR = M1 + 3
    _section(ws, R_HDR - 1, "B. Aturan transmisi (§5.2) — dari config/stress_parameters.yaml")
    _header(ws, R_HDR, ["kri_id", "kri_name", "aturan", "parameter"])
    names = ra.set_index("kri_id").kri_name
    r = R_HDR + 1
    for kid, rule in params["transmission"].items():
        text, par = describe_rule(kid, rule)
        for j, v in enumerate([kid, names[kid], text, par]):
            _put(ws, r, j + 1, v, wrap=j in (1, 2, 3))
        r += 1
    _put(ws, r, 1, "Lainnya")
    _put(ws, r, 2, ", ".join(params["unchanged"]), wrap=True)
    _put(ws, r, 3, "Tidak berubah (tidak sensitif makro dalam project ini; dinyatakan sebagai limitation)", wrap=True)
    _put(ws, r, 4, "—")
    ws.row_dimensions[r].height = 30

    srs = outputs["stress_risk_scores"].pivot(index="risk_type", columns="scenario", values="final_score")
    se = outputs["stress_enterprise_summary"].set_index("scenario")
    X_HDR = r + 3
    _section(ws, X_HDR - 1, "C. Hasil stress — final risk score (engine, presisi penuh; format 1 desimal)")
    _header(ws, X_HDR, ["risk_type", "baseline", "adverse", "severe", "Δ severe − baseline (formula)",
                        "bobot", "kontribusi severe (formula)"])
    X0 = X_HDR + 1
    for i, risk in enumerate(risks):
        rr = X0 + i
        _put(ws, rr, 1, risk)
        for j, s in enumerate(datastore.SCENARIOS):
            _put(ws, rr, 2 + j, float(srs.at[risk, s]), font=BLUE, fmt="0.0")
        _put(ws, rr, 5, f"=D{rr}-B{rr}", fmt="+0.00;-0.00;0.00")
        _put(ws, rr, 6, f"=INDEX({_q(S01)}!$F$6:$F$15,MATCH(A{rr},{_q(S01)}!$A$6:$A$15,0))", font=GREEN_LINK, fmt="0.00")
        _put(ws, rr, 7, f"=E{rr}*F{rr}", fmt="+0.000;-0.000;0.000")
    X1 = X0 + len(risks) - 1
    extra = [("Enterprise score", "enterprise_score", "0.0"), ("Enterprise rating", "rating", None),
             ("Limit utilization (%)", "limit_utilization", "0"), ("KRI skor ≥ 4", "tolerance_breaches", "0"),
             ("KRI skor = 5", "limit_breaches", "0")]
    for i, (label, col, fmt) in enumerate(extra):
        rr = X1 + 1 + i
        _put(ws, rr, 1, label, font=BOLD)
        for j, s in enumerate(datastore.SCENARIOS):
            v = se.at[s, col]
            _put(ws, rr, 2 + j, v.item() if hasattr(v, "item") else v, font=BLUE, fmt=fmt)
    _put(ws, X1 + 1, 7, f"=SUM(G{X0}:G{X1})", fmt="+0.000;-0.000;0.000")
    ENT_STRESS_ROW = X1 + 1

    c = params["capital_overlay"]
    P_HDR = X1 + len(extra) + 3
    _section(ws, P_HDR - 1, "D. Capital Adequacy Overlay (§5.5) — parameter (input) dan formula hidup")
    _header(ws, P_HDR, ["parameter", "nilai", "sumber"])
    par_rows = [("Pembiayaan bruto (Rp miliar)", c["credit_loss"]["financing"]),
                ("Loss rate kredit (asumsi tunggal, bukan model kerugian kredit)", c["credit_loss"]["loss_rate"]),
                ("Portofolio sukuk FVOCI", c["investment_loss"]["fvoci_portfolio"]),
                ("Kerugian operasional tahunan (basis)", c["operational_loss"]["annual_loss_base"]),
                ("Deposito mudharabah", c["dcr_cost"]["deposito_mudharabah"]),
                ("Pass-through biaya DCR", c["dcr_cost"]["pass_through"]),
                ("Modal KPMM awal", c["capital"]["capital_base"]), ("ATMR awal", c["atmr"]["atmr_base"])]
    P0 = P_HDR + 1
    for i, (label, v) in enumerate(par_rows):
        _put(ws, P0 + i, 1, label)
        _put(ws, P0 + i, 2, float(v), font=BLUE, fmt="#,##0.00", fill=INPUT_FILL)
        _put(ws, P0 + i, 3, "config/stress_parameters.yaml")
    p = {k: f"$B${P0 + i}" for i, k in enumerate(["fin", "lr", "fv", "op", "dep", "pt", "cap", "atmr"])}

    Q_HDR = P0 + len(par_rows) + 1
    _header(ws, Q_HDR, ["item", "baseline", "adverse", "severe"])
    Q0 = Q_HDR + 1
    st = outputs["stress_results"].pivot(index="kri_id", columns="scenario", values="stressed_value")
    base_cr01 = float(outputs["stress_results"].query("kri_id == 'CR01' and scenario == 'baseline'").base_value.iloc[0])
    cap_rows = ["ΔNPF (pp) = CR01 stressed − CR01 baseline", "Kerugian kredit", "Kerugian sukuk FVOCI",
                "Kerugian operasional", "Biaya displaced commercial risk", "Total kerugian", "Modal stress", "ATMR stress",
                "KPMM (%)", "Skor CAP01 (16/14/12/10,5)", "KPMM engine (data/output)", "Selisih formula − engine (pp)"]
    capo = outputs["capital_overlay"].set_index("scenario")
    for i, label in enumerate(cap_rows):
        _put(ws, Q0 + i, 1, label, font=BOLD if i in (5, 8) else BLACK)
    for j, s in enumerate(datastore.SCENARIOS):
        col = get_column_letter(2 + j)
        mcol = get_column_letter(3 + j)  # kolom skenario pada tabel makro
        rr = Q0
        _put(ws, rr, 2 + j, f"={float(st.at['CR01', s])!r}-{base_cr01!r}", fmt="0.00")
        _put(ws, rr + 1, 2 + j, f"={col}{rr}/100*{p['fin']}*{p['lr']}", fmt="#,##0")
        _put(ws, rr + 2, 2 + j, f"={p['fv']}*ABS({mac('sukuk_mv_shock_fvoci', mcol)})/100", fmt="#,##0")
        _put(ws, rr + 3, 2 + j, f"={p['op']}*{mac('operational_event_uplift', mcol)}/100", fmt="#,##0")
        _put(ws, rr + 4, 2 + j, f"={p['dep']}*{mac('policy_rate_shift', mcol)}/10000*{p['pt']}", fmt="#,##0")
        _put(ws, rr + 5, 2 + j, f"=SUM({col}{rr + 1}:{col}{rr + 4})", fmt="#,##0")
        _put(ws, rr + 6, 2 + j, f"={p['cap']}-{col}{rr + 5}", fmt="#,##0")
        _put(ws, rr + 7, 2 + j, f"={p['atmr']}*(1+{mac('rwa_uplift', mcol)}/100)", fmt="#,##0")
        _put(ws, rr + 8, 2 + j, f"={col}{rr + 6}/{col}{rr + 7}*100", fmt="0.0")
        cutc = [f"INDEX({TH}!${x}${T0}:${x}${T1},MATCH(\"CAP01\",{TH}!$A${T0}:$A${T1},0))" for x in "FGHI"]
        _put(ws, rr + 9, 2 + j, f'=IF({col}{rr + 8}>={cutc[0]},1,IF({col}{rr + 8}>={cutc[1]},2,IF({col}{rr + 8}>={cutc[2]},3,IF({col}{rr + 8}>={cutc[3]},4,5))))')
        _put(ws, rr + 10, 2 + j, float(capo.at[s, "kpmm"]), font=BLUE, fmt="0.0000")
        _put(ws, rr + 11, 2 + j, f"={col}{rr + 8}-{col}{rr + 10}", fmt="0.0000")
    K_ROW = {"kpmm": Q0 + 8, "total": Q0 + 5, "score": Q0 + 9}
    ws.cell(row=Q0, column=2).comment = Comment("ΔNPF diambil dari stress_results.csv (engine).", "ERM engine")

    W_HDR = Q0 + len(cap_rows) + 2
    _section(ws, W_HDR - 1, "E. Waterfall enterprise baseline → severe (§5.4, engine)")
    _header(ws, W_HDR, ["langkah", "driver", "kontribusi", "kumulatif"])
    wf = outputs["stress_waterfall"].query("to_scenario == 'severe'")
    for i, row in enumerate(wf.itertuples(index=False)):
        rr = W_HDR + 1 + i
        _put(ws, rr, 1, int(row.step))
        _put(ws, rr, 2, row.driver)
        _put(ws, rr, 3, float(row.contribution), font=BLUE, fmt="+0.000;-0.000;0.000")
        _put(ws, rr, 4, float(row.cumulative), font=BLUE, fmt="0.000")

    K_HDR = W_HDR + len(wf) + 3
    _section(ws, K_HDR - 1, "F. Nilai & skor KRI hasil stress (engine)")
    _header(ws, K_HDR, ["kri_id", "risk_type", "base value", "base score", "baseline value", "baseline score",
                        "adverse value", "adverse score", "severe value", "severe score"])
    sr = outputs["stress_results"]
    for i, kid in enumerate(ra.kri_id):
        rr = K_HDR + 1 + i
        g = sr[sr.kri_id == kid].set_index("scenario")
        vals = [kid, g.risk_type.iloc[0], float(g.base_value.iloc[0]), int(g.base_score.iloc[0])]
        for s in datastore.SCENARIOS:
            vals += [float(g.at[s, "stressed_value"]), int(g.at[s, "stressed_score"])]
        for j, v in enumerate(vals):
            _put(ws, rr, j + 1, v, font=BLUE if j >= 2 else BLACK, fmt="0.0000" if isinstance(v, float) else None)
    _widths(ws, {"A": 44, "B": 22, "C": 44, "D": 44, "E": 16, "F": 10, "G": 16, "H": 12, "I": 12, "J": 12})

    # ---- 08_Management_Actions
    ws = sheets[S08]
    _title(ws, "08 · Management Action Tracker (§8.4) + Emerging Risk Register (§7)",
           "Aturan: KRI skor ≥ 4 → action (skor 4: 30 hari; skor 5: 14 hari + eskalasi Direksi/RMC). Teks dari action_library.csv.")
    act = outputs["management_actions"]
    _header(ws, 4, ["status", "jumlah (formula)"])
    statuses = ["Open", "Overdue", "Closed", "Superseded (eskalasi skor 5)"]
    for i, stt in enumerate(statuses):
        _put(ws, 5 + i, 1, stt)
        _put(ws, 5 + i, 2, f'=COUNTIF($J$12:$J${12 + max(len(act), 1) - 1},A{5 + i})')
    _put(ws, 9, 1, "Total", font=BOLD)
    _put(ws, 9, 2, "=SUM(B5:B8)")
    _header(ws, 11, list(act.columns))
    for i, row in enumerate(act.itertuples(index=False)):
        for j, v in enumerate(row):
            _put(ws, 12 + i, j + 1, v, wrap=j in (4, 6), font=BLUE if j in (4,) else BLACK)
    er = ref["emerging_risks"]
    _section(ws, er_start, "Emerging risk register (data/reference/emerging_risks.csv)")
    _header(ws, er_start + 1, list(er.columns) + ["priority (formula)"])
    for i, row in enumerate(er.itertuples(index=False)):
        rr = er_start + 2 + i
        for j, v in enumerate(row):
            v = v.item() if hasattr(v, "item") else v
            _put(ws, rr, j + 1, v, wrap=j in (3, 9))
        _put(ws, rr, len(er.columns) + 1, f"=E{rr}*F{rr}*G{rr}/5", fmt="0.0")
    _widths(ws, {"A": 12, "B": 22, "C": 16, "D": 40, "E": 48, "F": 18, "G": 48, "H": 26, "I": 12, "J": 14, "K": 40, "L": 22, "M": 12, "N": 12})

    # ---- 09_Regulatory_Mapping
    ws = sheets[S09]
    _title(ws, f"09 · Regulatory Mapping — status diverifikasi per {REG_VERIFIED_ON} (§11)",
           "Status peraturan ditelusuri dari JDIH BPK dan laman regulasi OJK; lihat sumber di bawah tabel.")
    _write_reg_table(ws)

    # ---- 10_Change_Log
    ws = sheets[S10]
    _title(ws, "10 · Change Log — versi threshold & keputusan metodologi (approver fiktif)")
    _header(ws, 4, ["versi", "tanggal", "perubahan", "approver", "referensi"])
    for i, row in enumerate(CHANGE_LOG):
        for j, v in enumerate(row):
            _put(ws, 5 + i, j + 1, v, wrap=j == 2)
    _widths(ws, {"A": 8, "B": 12, "C": 70, "D": 34, "E": 24})

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
        w.page_setup.orientation = "landscape"
        w.page_setup.paperSize = w.PAPERSIZE_A4
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb._erm_cells = dict(enterprise_row=ENT_ROW, enterprise_display_row=ENT_DISPLAY_ROW, risk_rows=(B0, B1),
                         kri_rows=(A0, A1), credit_final_row=idx["fi"], capital_rows=K_ROW,
                         enterprise_stress_row=ENT_STRESS_ROW, c0=C0, emerging_start=er_start)
    return wb


def _find_soffice() -> str | None:
    import shutil
    for name in ("soffice", "libreoffice", "soffice.exe"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
                      r"C:\Program Files\LibreOffice\program\soffice.exe",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(candidate).exists():
            return candidate
    return None


def recalculate(path: Path, timeout: int = 180) -> dict:
    """Hitung ulang formula agar nilai tersimpan di file (dibutuhkan test & previewer).

    1) Skrip recalc (ERM_RECALC_SCRIPT) bila tersedia; 2) LibreOffice headless (convert xlsx → xlsx);
    3) tanpa keduanya: dilewati — Excel tetap menghitung ulang saat file dibuka."""
    path = Path(path)
    if RECALC_SCRIPT.exists():
        res = subprocess.run(["python3", str(RECALC_SCRIPT), str(path), str(timeout)], cwd=RECALC_SCRIPT.parent,
                             capture_output=True, text=True)
        try:
            return json.loads(res.stdout)
        except json.JSONDecodeError:
            pass
    soffice = _find_soffice()
    if soffice is None:
        return {"status": "skipped", "reason": "LibreOffice tidak ditemukan; nilai formula dihitung saat file dibuka di Excel"}
    import shutil
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp) / "profile"
        cmd = [soffice, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--calc",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", tmp, str(path)]
        try:
            subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            return {"error": "LibreOffice timeout"}
        converted = Path(tmp) / path.name
        if not converted.exists():
            return {"error": "LibreOffice gagal mengonversi file"}
        shutil.copyfile(converted, path)
    return {"status": "success", "method": "libreoffice-convert"}


REGULATORY_PATH = datastore.ROOT / "regulatory" / "Regulatory_Mapping.xlsx"


def build_regulatory_mapping(path: Path | str = REGULATORY_PATH) -> Path:
    """regulatory/Regulatory_Mapping.xlsx — rujukan §11 dengan status hasil verifikasi."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Regulatory_Mapping"
    _title(ws, f"Regulatory Mapping — status diverifikasi per {REG_VERIFIED_ON} (spec §11)",
           "Status peraturan ditelusuri dari JDIH BPK dan laman regulasi OJK; lihat sumber di bawah tabel.")
    end = _write_reg_table(ws)
    ws.cell(row=end + 2, column=1, value=DISCLAIMER).font = NOTE_FONT
    ws.sheet_view.showGridLines = False
    wb.save(path)
    return path


def build(path: Path | str = EXCEL_PATH, recalc: bool = True) -> dict:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    layout = wb._erm_cells
    wb.save(path)
    result = recalculate(path) if recalc else {"status": "skipped"}
    build_regulatory_mapping()
    (path.parent / "ERM_Framework_layout.json").write_text(json.dumps(layout, indent=2), encoding="utf-8")
    return result
